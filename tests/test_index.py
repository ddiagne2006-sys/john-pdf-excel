"""
Tests de l'index SQLite. Tout se passe dans un dossier temporaire : ni PDF_Colis,
ni le vrai %LOCALAPPDATA% ne sont utilisés.

Lancement :  python -m unittest discover tests
"""

from pathlib import Path
from unittest import mock
import contextlib
import io
import os
import sqlite3
import sys
import tempfile
import time
import unittest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import pymupdf  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

import main  # noqa: E402

NUMERO = "COLIS-2026-00481"
TESSDATA = main.trouver_tessdata()


def lignes_bordereau(reference, montant="850 000"):
    return ["BORDEREAU D'EXPEDITION", f"Ref : {reference}", "Date : 15/10/2026",
            f"Montant : {montant} GNF", "Poids : 25,5 kg", "Nombre de colis : 4"]


def creer_pdf_texte(chemin, lignes):
    chemin.parent.mkdir(parents=True, exist_ok=True)
    with pymupdf.open() as document:
        page = document.new_page()
        for i, ligne in enumerate(lignes):
            page.insert_text((72, 72 + 20 * i), ligne, fontsize=12)
        document.save(str(chemin))


def creer_pdf_scan(chemin, lignes):
    """Page texte rendue en image puis remise seule dans un PDF : aucune couche texte."""
    with pymupdf.open() as source:
        page = source.new_page()
        for i, ligne in enumerate(lignes):
            page.insert_text((72, 72 + 20 * i), ligne, fontsize=12)
        image = page.get_pixmap(dpi=200)
    with pymupdf.open() as document:
        page = document.new_page()
        page.insert_image(page.rect, pixmap=image)
        document.save(str(chemin))


def creer_pdf_protege(chemin):
    with pymupdf.open() as document:
        document.new_page().insert_text((72, 72), f"Ref : {NUMERO}")
        document.save(str(chemin), encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw="o", user_pw="u")


def toucher(chemin):
    """Avance la date de modification d'une seconde (sûr même si l'écriture a été très rapide)."""
    etat = chemin.stat()
    os.utime(chemin, ns=(etat.st_atime_ns, etat.st_mtime_ns + 1_000_000_000))


class TestIndex(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        racine = Path(temp.name)
        self.dossier = racine / "PDF client"          # espace dans le nom, comme chez un client
        self.dossier_index = racine / "index"
        self.sorties = racine / "sorties"
        self.sorties.mkdir()

        # Sécurité : même un appel sans dossier_index n'écrit pas dans le vrai %LOCALAPPDATA%
        patch = mock.patch.object(main, "DOSSIER_INDEX", racine / "index_par_defaut")
        patch.start()
        self.addCleanup(patch.stop)

        creer_pdf_texte(self.dossier / "texte_00481.pdf", lignes_bordereau(NUMERO))
        creer_pdf_texte(self.dossier / "texte_00482.pdf", lignes_bordereau("COLIS-2026-00482"))
        creer_pdf_texte(self.dossier / "2026" / "facture.pdf",
                        ["FACTURE F-118", "Objet : colis 2026 00481", "Total : 1 250 000,50 GNF"])
        creer_pdf_scan(self.dossier / "scan_00481.pdf", lignes_bordereau(NUMERO))
        creer_pdf_protege(self.dossier / "protege.pdf")
        (self.dossier / "abime.pdf").write_bytes(b"%PDF-1.7\n" + os.urandom(2000))
        self.nb_pdf = 6

    # -- outils ----------------------------------------------------------------------
    def lancer(self, avec_index=True, tessdata=TESSDATA, reconstruire=False):
        """Recherche ; renvoie (resultats, non_analyses, statistiques, texte affiché, durée)."""
        sortie = io.StringIO()
        debut = time.perf_counter()
        with contextlib.redirect_stdout(sortie):
            resultats, non_analyses, _, statistiques = main.rechercher(
                NUMERO, self.dossier, tessdata, avec_index=avec_index,
                reconstruire=reconstruire, dossier_index=self.dossier_index)
        return resultats, non_analyses, statistiques, sortie.getvalue(), time.perf_counter() - debut

    def feuilles(self, resultats, non_analyses):
        """Contenu des deux feuilles Excel produites (valeurs de toutes les cellules)."""
        fichier = self.sorties / f"resultat_{len(list(self.sorties.iterdir()))}.xlsx"
        main.exporter_excel(NUMERO, resultats, non_analyses, self.dossier, fichier)
        classeur = load_workbook(fichier)
        return {feuille.title: list(feuille.iter_rows(values_only=True)) for feuille in classeur}

    def fichiers_trouves(self, resultats):
        return sorted(str(r["fichier"].relative_to(self.dossier)) for r in resultats)

    def fichier_index(self):
        return main.chemin_index(self.dossier, self.dossier_index)

    # -- tests -----------------------------------------------------------------------
    def test_resultats_identiques_a_sans_index(self):
        direct = self.feuilles(*self.lancer(avec_index=False)[:2])
        premier = self.feuilles(*self.lancer()[:2])      # index construit
        deuxieme = self.feuilles(*self.lancer()[:2])     # tout depuis le cache
        self.assertEqual(premier, direct)
        self.assertEqual(deuxieme, direct)
        # Les deux PDF illisibles sont bien dans « Non analysés »
        non_analyses = [ligne[0] for ligne in direct["Non analysés"][1:]]
        self.assertIn("protege.pdf", non_analyses)
        self.assertIn("abime.pdf", non_analyses)

    def test_deuxieme_lancement_depuis_le_cache(self):
        *_, stats1, _, duree1 = self.lancer()
        *_, stats2, _, duree2 = self.lancer()
        self.assertEqual((stats1["lus"], stats1["cache"]), (self.nb_pdf, 0))
        self.assertEqual((stats2["lus"], stats2["cache"], stats2["retires"]), (0, self.nb_pdf, 0))
        self.assertLess(duree2, duree1)
        sys.stderr.write(f"\n  [temps] 1er lancement {duree1:.3f} s, 2e (cache) {duree2:.3f} s "
                         f"({duree1 / duree2:.0f}x plus rapide) ")

    def test_pdf_modifie_seul_relu(self):
        self.lancer()
        chemin = self.dossier / "texte_00481.pdf"
        creer_pdf_texte(chemin, lignes_bordereau(NUMERO, montant="999 000"))
        toucher(chemin)
        resultats, _, stats, *_ = self.lancer()
        self.assertEqual((stats["lus"], stats["cache"]), (1, self.nb_pdf - 1))
        nombres = next(r["nombres"] for r in resultats if r["fichier"] == chemin)
        self.assertIn("999 000", nombres)
        self.assertNotIn("850 000", nombres)

    def test_pdf_supprime_et_nouveau(self):
        self.lancer()
        (self.dossier / "texte_00481.pdf").unlink()
        creer_pdf_texte(self.dossier / "nouveau_00481.pdf", lignes_bordereau(NUMERO))
        resultats, _, stats, *_ = self.lancer()
        self.assertEqual((stats["lus"], stats["retires"]), (1, 1))
        trouves = self.fichiers_trouves(resultats)
        self.assertNotIn("texte_00481.pdf", trouves)
        self.assertIn("nouveau_00481.pdf", trouves)

    @unittest.skipUnless(TESSDATA, "Tesseract (fra+eng) introuvable sur cette machine")
    def test_scans_lus_quand_tesseract_arrive(self):
        resultats, non_analyses, stats1, *_ = self.lancer(tessdata=None)
        self.assertNotIn("scan_00481.pdf", self.fichiers_trouves(resultats))
        self.assertIn((self.dossier / "scan_00481.pdf", f"{main.RAISON_SANS_OCR} (page(s) 1)"), non_analyses)

        resultats, non_analyses, stats2, *_ = self.lancer(tessdata=TESSDATA)
        self.assertEqual(stats2["lus"], 1)   # seul le scan est relu
        scan = next(r for r in resultats if r["fichier"].name == "scan_00481.pdf")
        self.assertEqual(scan["lecture"], "OCR")
        self.assertIn("850 000", scan["nombres"])
        self.assertFalse(any(f.name == "scan_00481.pdf" for f, _ in non_analyses))

    def test_pdf_illisible_garde_en_cache_jusqu_a_modification(self):
        self.lancer()
        _, non_analyses, stats, *_ = self.lancer()
        self.assertEqual(stats["lus"], 0)
        self.assertEqual({f.name for f, _ in non_analyses}, {"protege.pdf", "abime.pdf"})

        abime = self.dossier / "abime.pdf"
        abime.write_bytes(b"%PDF-1.7\n" + os.urandom(3000))
        toucher(abime)
        *_, stats, _, _ = self.lancer()
        self.assertEqual(stats["lus"], 1)

    def test_index_corrompu_reconstruit(self):
        reference = self.feuilles(*self.lancer()[:2])
        self.fichier_index().write_bytes(os.urandom(8192))
        resultats, non_analyses, stats, affichage, _ = self.lancer()
        self.assertIn("index illisible", affichage)
        self.assertEqual(stats["lus"], self.nb_pdf)
        self.assertEqual(self.feuilles(resultats, non_analyses), reference)
        # Et l'index reconstruit est de nouveau utilisable
        *_, stats, _, _ = self.lancer()
        self.assertEqual(stats["lus"], 0)

    def test_changement_ocr_dpi_reconstruit(self):
        self.lancer()
        with mock.patch.object(main, "OCR_DPI", 200):
            *_, stats, affichage, _ = self.lancer()
        self.assertEqual(stats["lus"], self.nb_pdf)
        self.assertIn("réglages de lecture modifiés", affichage)

    def test_reindexer(self):
        self.lancer()
        *_, stats, affichage, _ = self.lancer(reconstruire=True)
        self.assertEqual(stats["lus"], self.nb_pdf)
        self.assertIn("reconstruction complète demandée", affichage)

    def test_index_hors_du_dossier_des_pdf_et_en_wal(self):
        self.lancer()
        self.assertEqual(list(self.dossier.rglob("*.sqlite*")), [])
        fichier = self.fichier_index()
        self.assertEqual(fichier.parent, self.dossier_index)
        self.assertRegex(fichier.name, r"^index_[0-9a-f]{16}\.sqlite$")
        with contextlib.closing(sqlite3.connect(fichier)) as connexion:
            self.assertEqual(connexion.execute("PRAGMA journal_mode").fetchone()[0], "wal")

    def test_un_index_par_dossier(self):
        autre = self.dossier.parent / "Autre client"
        self.assertNotEqual(main.chemin_index(self.dossier, self.dossier_index),
                            main.chemin_index(autre, self.dossier_index))

    def test_indexer_seulement_sans_numero(self):
        arguments = ["main.py", "--indexer-seulement", "--dossier", str(self.dossier)]
        sortie = io.StringIO()
        with mock.patch.object(sys, "argv", arguments), \
             mock.patch.object(main, "DOSSIER_INDEX", self.dossier_index), \
             mock.patch("builtins.input", side_effect=AssertionError("numéro demandé")), \
             contextlib.redirect_stdout(sortie):
            code = main.main()
        self.assertEqual(code, 0)
        self.assertIn(f"Index : {self.nb_pdf} PDF lus, 0 depuis le cache, 0 retirés", sortie.getvalue())

    def test_lien_par_formule_hyperlink(self):
        resultats, non_analyses, *_ = self.lancer()
        fichier = self.sorties / "liens.xlsx"
        main.exporter_excel(NUMERO, resultats, non_analyses, self.dossier, fichier)
        feuille = load_workbook(fichier)["Résultats"]
        self.assertEqual(feuille._hyperlinks, [])     # plus aucun lien hypertexte classique
        attendus = sorted('=HYPERLINK("{}","{}")'.format(r["fichier"], r["fichier"].relative_to(self.dossier))
                          for r in resultats)
        cellules = feuille["A"][3:]
        self.assertEqual(sorted(c.value for c in cellules), attendus)
        self.assertTrue(all(c.style == "Hyperlink" for c in cellules))

    def test_lien_guillemets_doubles_et_chemin_trop_long(self):
        court = {"fichier": self.dossier / 'a"b.pdf', "pages": [1], "references": [NUMERO],
                 "nombres": [], "lecture": "texte"}
        long_ = dict(court, fichier=self.dossier / ("x" * 260 + ".pdf"))
        fichier = self.sorties / "limites.xlsx"
        main.exporter_excel(NUMERO, [court, long_], [], self.dossier, fichier)
        feuille = load_workbook(fichier)["Résultats"]
        chemin = str(court["fichier"]).replace('"', '""')
        self.assertEqual(feuille["A4"].value, f'=HYPERLINK("{chemin}","a""b.pdf")')
        self.assertEqual(feuille["A5"].value, "x" * 260 + ".pdf")   # nom seul, sans lien
        self.assertEqual(feuille["A5"].style, "Hyperlink")

    def test_sans_index_incompatible_avec_reindexer(self):
        with mock.patch.object(sys, "argv", ["main.py", NUMERO, "--sans-index", "--reindexer"]), \
             contextlib.redirect_stderr(io.StringIO()), \
             self.assertRaises(SystemExit) as sortie:
            main.lire_arguments()
        self.assertEqual(sortie.exception.code, 1)


if __name__ == "__main__":
    unittest.main()
