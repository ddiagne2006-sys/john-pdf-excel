"""
DDT PDF Excel - étape 1 : moteur de recherche dans un dossier de PDF.

Recherche une référence (ex. COLIS-2026-00481) dans tous les PDF du dossier
PDF_Colis (sous-dossiers compris) et écrit les nombres trouvés dans resultat.xlsx.

Prérequis :  pip install pymupdf openpyxl
OCR       :  Tesseract avec les langues fra et eng (facultatif : sans lui, les scans sont
             listés dans « Non analysés »)
Usage     :  python main.py COLIS-2026-00481 [--dossier DOSSIER] [--sortie FICHIER.xlsx]
Version .exe (build.ps1) : moteur\\Recherche_PDF.exe, mêmes arguments ; sans argument
             (double-clic), le numéro et le dossier sont demandés dans la fenêtre.
"""

from pathlib import Path
import argparse
import hashlib
import json
import os
import re
import sqlite3
import sys
import time

import pymupdf  # PyMuPDF >= 1.24.3 (avant : "import fitz")
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

# ----------------------------------------------------------------------------
# Configuration
# ----------------------------------------------------------------------------
if getattr(sys, "frozen", False):
    # Version .exe (PyInstaller) : ...\Recherche_PDF\moteur\Recherche_PDF.exe
    BASE_DIR = Path(sys.executable).resolve().parent    # moteur\ (tesseract\tessdata)
    DOSSIER_OUTIL = BASE_DIR.parent                     # dossier livré (classeur, PDF_Colis)
else:
    BASE_DIR = DOSSIER_OUTIL = Path(__file__).resolve().parent
PDF_DIR = DOSSIER_OUTIL / "PDF_Colis"
EXCEL_FILE = DOSSIER_OUTIL / "resultat.xlsx"

# "page"     : nombres de la (des) page(s) où la référence apparaît
# "document" : nombres de tout le PDF
PORTEE = "page"

# Limite d'une cellule Excel : 32 767 caractères
LIMITE_CELLULE = 32_000

SEPARATEURS = r"[\s\-_/.]*"

# OCR des pages scannées (Tesseract, via PyMuPDF)
OCR_LANGUES = "fra+eng"
OCR_DPI = 300
# En dessous de ce nombre de caractères non blancs, une page est considérée comme scannée
SEUIL_TEXTE = 20
# Scan d'appli téléphone : image sur toute la page + filigrane texte (« Scanned with CamScanner »).
# Moins de SEUIL_TEXTE_FILIGRANE caractères ET images sur au moins SEUIL_COUVERTURE_IMAGES de la page.
SEUIL_TEXTE_FILIGRANE = 200
SEUIL_COUVERTURE_IMAGES = 0.70
# Installations standard de Tesseract sous Windows
CHEMINS_TESSDATA = [
    Path(os.environ.get("ProgramFiles", r"C:\Program Files")) / "Tesseract-OCR" / "tessdata",
    Path(os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)")) / "Tesseract-OCR" / "tessdata",
    Path(os.environ.get("LOCALAPPDATA", Path.home() / "AppData" / "Local")) / "Programs" / "Tesseract-OCR" / "tessdata",
]
RAISON_SANS_OCR = "OCR non disponible (Tesseract introuvable)"
RAISON_VIDE_APRES_OCR = "Aucun texte, même après OCR"
PREFIXE_ERREUR_OCR = "Erreur OCR"

# Index des textes déjà lus (SQLite), un fichier par dossier de PDF.
# Jamais dans le dossier des PDF : il peut être sur OneDrive, sur un réseau ou en lecture seule.
DOSSIER_INDEX = Path(os.environ.get("LOCALAPPDATA", Path.home() / "AppData" / "Local")) / "DDT_Recherche_PDF"
# À augmenter quand le contenu de l'index change de forme : l'index est alors reconstruit
INDEX_VERSION = 1
PROGRESSION_TOUS_LES = 10   # PDF


# ----------------------------------------------------------------------------
# Recherche de la référence
# ----------------------------------------------------------------------------
def motif_reference(numero):
    """
    Construit une expression qui trouve la référence même si la ponctuation
    ou les espaces diffèrent (COLIS-2026-00481 = colis 2026 00481 = COLIS/2026/00481),
    mais sans accepter une référence plus longue (COLIS-2026-004810)
    ni une partie d'un autre mot.
    """
    morceaux = re.findall(r"[A-Za-z0-9]+", numero)
    if not morceaux:
        raise ValueError("Le numéro doit contenir des lettres ou des chiffres.")
    corps = SEPARATEURS.join(re.escape(m) for m in morceaux)
    return re.compile(rf"(?<![A-Za-z0-9]){corps}(?![A-Za-z0-9])", re.IGNORECASE)


# ----------------------------------------------------------------------------
# Extraction des nombres
# ----------------------------------------------------------------------------
# Unités qui peuvent être collées à un nombre (850000GNF, 25,5kg) : seul le nombre est gardé.
# Insensibles à la casse.
UNITES = ["GNF", "FG", "FCFA", "XOF", "USD", "EUR", "kg", "g", "t", "m3", "cm", "pcs"]

# Un nombre "français" : 850 000 / 850.000 / 25,5 / 1 250 000,75
# ou une référence contenant des chiffres : COLIS-2026-00481, 15/10/2026, +224622123456
MOTIF_NOMBRE = re.compile(
    r"""
    (?<![A-Za-z0-9])
    (?:
        (?P<nombre>\d{1,3}(?:[   ]\d{3})+(?:,\d+)?|\d+(?:,\d+)?)
        (?i:""" + "|".join(sorted(map(re.escape, UNITES), key=len, reverse=True)) + r""")
                                                # nombre + unité collée : 850000GNF / 25,5kg
      | \d{1,3}(?:[   ]\d{3})+(?:,\d+)?        # milliers : 850 000 / 1 250 000,50
      | \+?[A-Za-z0-9]+(?:[\-/_.][A-Za-z0-9]+)*(?:,\d+)?  # 25 / 25,5 / 15/10/2026 / COLIS-2026-00481
    )
    (?![A-Za-z0-9])
    """,
    re.VERBOSE,
)


def extraire_nombres(texte):
    """Récupère, dans l'ordre et sans doublon, les éléments contenant des chiffres."""
    vus = set()
    resultat = []
    for m in MOTIF_NOMBRE.finditer(texte):
        # Nombre suivi d'une unité collée : on ne garde que le nombre
        valeur = m.group("nombre") or m.group()
        valeur = re.sub(r"[  ]", " ", valeur)
        if not re.search(r"\d", valeur):   # mots sans chiffre : ignorés
            continue
        if valeur not in vus:
            vus.add(valeur)
            resultat.append(valeur)
    return resultat


# ----------------------------------------------------------------------------
# Lecture des PDF
# ----------------------------------------------------------------------------
def lister_pdf(dossier):
    """Tous les PDF du dossier et de ses sous-dossiers, quelle que soit la casse (.pdf/.PDF)."""
    return sorted(p for p in dossier.rglob("*") if p.is_file() and p.suffix.lower() == ".pdf")


def trouver_tessdata():
    """
    Dossier tessdata qui contient toutes les langues de OCR_LANGUES, ou None.
    Ordre : variable TESSDATA_PREFIX, tesseract\\tessdata à côté de main.py (version .exe),
    puis installations standard. Un dossier absent ou incomplet est ignoré.
    """
    candidats = []
    prefixe = os.environ.get("TESSDATA_PREFIX")
    if prefixe:
        # Selon les versions de Tesseract, la variable vise tessdata ou son dossier parent
        candidats += [Path(prefixe), Path(prefixe) / "tessdata"]
    candidats.append(BASE_DIR / "tesseract" / "tessdata")
    candidats += CHEMINS_TESSDATA

    langues = OCR_LANGUES.split("+")
    for dossier in candidats:
        if all((dossier / f"{langue}.traineddata").is_file() for langue in langues):
            return dossier
    return None


def couverture_images(page):
    """Part de la page (de 0 à 1) couverte par des images : union des zones, bornée à la page."""
    zone = page.rect
    rects = [pymupdf.Rect(info["bbox"]) & zone for info in page.get_image_info()]
    rects = [r for r in rects if not r.is_empty]
    if not rects or zone.is_empty:
        return 0.0
    # Union exacte (des images qui se chevauchent ne comptent qu'une fois) :
    # bandes verticales découpées aux bords des images, hauteur couverte dans chaque bande
    bords = sorted({x for r in rects for x in (r.x0, r.x1)})
    surface = 0.0
    for gauche, droite in zip(bords, bords[1:]):
        couvert, bas = 0.0, float("-inf")
        for haut, fin in sorted((r.y0, r.y1) for r in rects if r.x0 <= gauche and r.x1 >= droite):
            if fin > bas:
                couvert += fin - max(haut, bas)
                bas = fin
        surface += (droite - gauche) * couvert
    return surface / (zone.width * zone.height)


def page_sans_texte(texte, page=None):
    """
    Vrai si la page doit passer à l'OCR :
    - moins de SEUIL_TEXTE caractères non blancs ;
    - ou, si la page est fournie, moins de SEUIL_TEXTE_FILIGRANE caractères et des images
      sur au moins SEUIL_COUVERTURE_IMAGES de sa surface (scan avec filigrane texte).
    """
    nb_caracteres = len("".join(texte.split()))
    if nb_caracteres < SEUIL_TEXTE:
        return True
    return (page is not None
            and nb_caracteres < SEUIL_TEXTE_FILIGRANE
            and couverture_images(page) >= SEUIL_COUVERTURE_IMAGES)


def lire_pages(fichier, tessdata=None):
    """
    Lit chaque page : son texte, ou l'OCR si elle n'en a pas (scan) et que tessdata est connu.
    Renvoie (pages, lectures, problemes, durees_ocr) :
      pages      : texte de chaque page
      lectures   : "texte" ou "OCR" pour chaque page
      problemes  : [(raison, numéro de page)] pour la feuille « Non analysés »
      durees_ocr : durée en secondes de chaque OCR
    """
    pages, lectures, problemes, durees_ocr = [], [], [], []
    with pymupdf.open(fichier) as pdf:
        if pdf.needs_pass:
            raise PermissionError("PDF protégé par mot de passe")
        for numero_page, page in enumerate(pdf, start=1):
            texte = page.get_text()
            lecture = "texte"
            if page_sans_texte(texte, page):
                if tessdata is None:
                    problemes.append((RAISON_SANS_OCR, numero_page))
                else:
                    debut = time.perf_counter()
                    try:
                        ocr = page.get_textpage_ocr(language=OCR_LANGUES, dpi=OCR_DPI,
                                                    full=True, tessdata=str(tessdata))
                        # Seul le texte OCR est gardé : le filigrane texte éventuel est remplacé
                        texte = page.get_text(textpage=ocr)
                        lecture = "OCR"
                        if page_sans_texte(texte):
                            problemes.append((RAISON_VIDE_APRES_OCR, numero_page))
                    except Exception as erreur:   # une page en échec n'empêche pas de lire les autres
                        problemes.append((f"{PREFIXE_ERREUR_OCR} : {erreur}", numero_page))
                    durees_ocr.append(time.perf_counter() - debut)
            pages.append(texte)
            lectures.append(lecture)
    return pages, lectures, problemes, durees_ocr


def lire_pdf(fichier, tessdata=None):
    """
    Lit un PDF et renvoie (document, durees_ocr). Le document est un dict :
    fichier, erreur (PDF illisible : mot de passe, fichier abîmé…), pages, lectures, problemes.
    """
    try:
        pages, lectures, problemes, durees_ocr = lire_pages(fichier, tessdata)
    except Exception as erreur:
        return {"fichier": fichier, "erreur": str(erreur), "pages": [], "lectures": [], "problemes": []}, []
    document = {"fichier": fichier, "erreur": None, "pages": pages, "lectures": lectures, "problemes": problemes}
    return document, durees_ocr


def probleme_provisoire(raison):
    """Problème qui peut disparaître au lancement suivant (Tesseract installé, erreur passagère)."""
    return raison == RAISON_SANS_OCR or raison.startswith(PREFIXE_ERREUR_OCR)


# ----------------------------------------------------------------------------
# Index : textes des PDF en cache (SQLite) pour ne relire que ce qui a changé
# ----------------------------------------------------------------------------
SCHEMA_INDEX = """
CREATE TABLE IF NOT EXISTS reglages (
    cle    TEXT PRIMARY KEY,
    valeur TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS pdf (
    chemin   TEXT PRIMARY KEY,          -- relatif au dossier des PDF
    taille   INTEGER NOT NULL,
    mtime_ns INTEGER NOT NULL,
    erreur   TEXT,                      -- PDF illisible (gardé tant que le fichier ne change pas)
    a_relire INTEGER NOT NULL           -- 1 : une page attend l'OCR, à relire au prochain lancement
);
CREATE TABLE IF NOT EXISTS page (
    chemin   TEXT NOT NULL,
    numero   INTEGER NOT NULL,
    texte    TEXT NOT NULL,
    lecture  TEXT NOT NULL,             -- "texte" ou "OCR"
    probleme TEXT,
    PRIMARY KEY (chemin, numero)
);
"""


def chemin_index(dossier, dossier_index=None):
    """Un index par dossier de PDF : index_<empreinte du chemin absolu>.sqlite."""
    empreinte = hashlib.sha256(os.path.normcase(str(dossier)).encode("utf-8")).hexdigest()[:16]
    return Path(dossier_index or DOSSIER_INDEX) / f"index_{empreinte}.sqlite"


def signature_reglages():
    """Réglages qui changent le texte extrait : s'ils changent, l'index est reconstruit."""
    return json.dumps({
        "index_version": INDEX_VERSION,
        "ocr_langues": OCR_LANGUES,
        "ocr_dpi": OCR_DPI,
        "seuil_texte": SEUIL_TEXTE,
        "seuil_texte_filigrane": SEUIL_TEXTE_FILIGRANE,
        "seuil_couverture_images": SEUIL_COUVERTURE_IMAGES,
    }, sort_keys=True)


def connecter_index(fichier):
    connexion = sqlite3.connect(fichier, timeout=30)   # attend jusqu'à 30 s un autre lancement
    try:
        connexion.execute("PRAGMA journal_mode=WAL")   # lectures et écriture simultanées
        connexion.execute("PRAGMA synchronous=NORMAL")
        if connexion.execute("PRAGMA quick_check").fetchone()[0] != "ok":
            raise sqlite3.DatabaseError("contrôle d'intégrité échoué")
        connexion.executescript(SCHEMA_INDEX)
    except sqlite3.Error:
        connexion.close()   # sinon Windows garde le fichier verrouillé et il ne peut pas être supprimé
        raise
    return connexion


def supprimer_index(fichier):
    for suffixe in ("", "-wal", "-shm"):
        Path(f"{fichier}{suffixe}").unlink(missing_ok=True)


def ouvrir_index(fichier, reconstruire=False):
    """
    Ouvre ou crée l'index. Il est vidé si on le demande ou si les réglages ont changé,
    supprimé puis recréé s'il est corrompu.
    """
    fichier.parent.mkdir(parents=True, exist_ok=True)
    try:
        connexion = connecter_index(fichier)
    except sqlite3.DatabaseError as erreur:
        print(f"Attention : index illisible ({erreur}). Il est supprimé et reconstruit.")
        supprimer_index(fichier)
        connexion = connecter_index(fichier)

    signature = signature_reglages()
    ligne = connexion.execute("SELECT valeur FROM reglages WHERE cle = 'signature'").fetchone()
    if reconstruire or ligne is None or ligne[0] != signature:
        if reconstruire:
            print("Index : reconstruction complète demandée.")
        elif ligne is not None:
            print("Index : réglages de lecture modifiés, reconstruction complète.")
        with connexion:
            connexion.execute("DELETE FROM page")
            connexion.execute("DELETE FROM pdf")
            connexion.execute("INSERT OR REPLACE INTO reglages VALUES ('signature', ?)", (signature,))
    return connexion


def enregistrer_document(connexion, relatif, etat, document):
    """Remplace dans l'index le PDF et ses pages (une transaction par PDF)."""
    problemes = {numero: raison for raison, numero in document["problemes"]}
    a_relire = any(probleme_provisoire(raison) for raison in problemes.values())
    with connexion:
        connexion.execute("DELETE FROM page WHERE chemin = ?", (relatif,))
        connexion.execute("INSERT OR REPLACE INTO pdf VALUES (?, ?, ?, ?, ?)",
                          (relatif, etat.st_size, etat.st_mtime_ns, document["erreur"], int(a_relire)))
        connexion.executemany(
            "INSERT INTO page VALUES (?, ?, ?, ?, ?)",
            [(relatif, numero, texte, lecture, problemes.get(numero))
             for numero, (texte, lecture) in enumerate(zip(document["pages"], document["lectures"]), start=1)],
        )


def charger_index(connexion, fichiers):
    """Documents de l'index, dans l'ordre de fichiers ({chemin relatif: chemin complet})."""
    erreurs = dict(connexion.execute("SELECT chemin, erreur FROM pdf"))
    pages = {}
    for chemin, texte, lecture, probleme, numero in connexion.execute(
            "SELECT chemin, texte, lecture, probleme, numero FROM page ORDER BY chemin, numero"):
        pages.setdefault(chemin, []).append((texte, lecture, probleme, numero))
    documents = []
    for relatif, fichier in fichiers.items():
        lignes = pages.get(relatif, [])
        documents.append({
            "fichier": fichier,
            "erreur": erreurs.get(relatif),
            "pages": [texte for texte, _, _, _ in lignes],
            "lectures": [lecture for _, lecture, _, _ in lignes],
            "problemes": [(probleme, numero) for _, _, probleme, numero in lignes if probleme],
        })
    return documents


def indexer(dossier, tessdata=None, reconstruire=False, dossier_index=None):
    """
    Met l'index à jour : relit les PDF nouveaux, modifiés (taille ou date) ou en attente d'OCR,
    retire les PDF supprimés. Renvoie (documents, durees_ocr, statistiques).
    """
    debut = time.perf_counter()
    connexion = ouvrir_index(chemin_index(dossier, dossier_index), reconstruire)
    try:
        en_cache = {chemin: (taille, mtime_ns, a_relire) for chemin, taille, mtime_ns, a_relire
                    in connexion.execute("SELECT chemin, taille, mtime_ns, a_relire FROM pdf")}

        fichiers, a_lire = {}, []
        for fichier in lister_pdf(dossier):
            try:
                etat = fichier.stat()
            except OSError:   # supprimé entre la liste et la lecture
                continue
            relatif = str(fichier.relative_to(dossier))
            fichiers[relatif] = fichier
            cache = en_cache.get(relatif)
            if cache is None or cache[:2] != (etat.st_size, etat.st_mtime_ns) or cache[2]:
                a_lire.append((relatif, fichier, etat))

        retires = [chemin for chemin in en_cache if chemin not in fichiers]
        with connexion:
            for chemin in retires:
                connexion.execute("DELETE FROM page WHERE chemin = ?", (chemin,))
                connexion.execute("DELETE FROM pdf WHERE chemin = ?", (chemin,))

        durees_ocr = []
        for i, (relatif, fichier, etat) in enumerate(a_lire, start=1):
            # Taille et date relevées AVANT la lecture : un PDF modifié pendant la lecture sera relu
            document, durees = lire_pdf(fichier, tessdata)
            durees_ocr += durees
            enregistrer_document(connexion, relatif, etat, document)
            if i % PROGRESSION_TOUS_LES == 0:
                print(f"{i}/{len(a_lire)} PDF…", flush=True)

        documents = charger_index(connexion, fichiers)
    finally:
        connexion.close()

    statistiques = {
        "lus": len(a_lire),
        "cache": len(fichiers) - len(a_lire),
        "retires": len(retires),
        "duree": time.perf_counter() - debut,
    }
    return documents, durees_ocr, statistiques


def charger_documents(dossier, tessdata=None, avec_index=True, reconstruire=False, dossier_index=None):
    """
    Textes de tous les PDF du dossier : depuis l'index, ou lus directement (avec_index=False,
    ou index inutilisable). Renvoie (documents, durees_ocr, statistiques de l'index ou None).
    """
    if not dossier.is_dir():
        raise FileNotFoundError(f"Le dossier PDF est introuvable : {dossier}")

    if avec_index:
        try:
            return indexer(dossier, tessdata, reconstruire, dossier_index)
        except (sqlite3.Error, OSError) as erreur:
            print(f"Attention : index inutilisable ({erreur}). Lecture directe de tous les PDF.")

    documents, durees_ocr = [], []
    for fichier in lister_pdf(dossier):
        document, durees = lire_pdf(fichier, tessdata)
        documents.append(document)
        durees_ocr += durees
    return documents, durees_ocr, None


# ----------------------------------------------------------------------------
# Recherche de la référence dans les textes
# ----------------------------------------------------------------------------
def chercher(motif, documents):
    """Renvoie (resultats, non_analyses) pour les documents lus (voir lire_pdf())."""
    resultats = []
    non_analyses = []   # (fichier, raison) : PDF illisibles, pages scannées non lues

    for document in documents:
        fichier = document["fichier"]
        if document["erreur"] is not None:
            non_analyses.append((fichier, document["erreur"]))
            continue
        pages, lectures, problemes = document["pages"], document["lectures"], document["problemes"]

        # Une ligne par raison, avec les pages concernées
        pages_par_raison = {}
        for raison, numero_page in problemes:
            pages_par_raison.setdefault(raison, []).append(str(numero_page))
        for raison, numeros in pages_par_raison.items():
            non_analyses.append((fichier, f"{raison} (page(s) {', '.join(numeros)})"))

        pages_trouvees = [i for i, t in enumerate(pages, start=1) if motif.search(t)]
        if not pages_trouvees:
            continue

        if PORTEE == "document":
            texte = "\n".join(pages)
        else:
            texte = "\n".join(pages[i - 1] for i in pages_trouvees)

        # Formes exactes de la référence dans le texte ("COLIS-2026-00481", "colis 2026 00481"…),
        # puis retrait de la référence pour qu'elle ne soit pas découpée dans les nombres.
        # Remplacée par un saut de ligne : il ne peut pas recoller deux nombres en "milliers".
        references = list(dict.fromkeys(m.group() for m in motif.finditer(texte)))
        texte = motif.sub("\n", texte)

        # Comment ont été lues les pages où la référence a été trouvée
        modes = {lectures[i - 1] for i in pages_trouvees}

        resultats.append({
            "fichier": fichier,
            "pages": pages_trouvees,
            "references": references,
            "nombres": extraire_nombres(texte),
            "lecture": " + ".join(m for m in ("texte", "OCR") if m in modes),
        })

    return resultats, non_analyses


def rechercher(numero, dossier=PDF_DIR, tessdata=None, avec_index=True, reconstruire=False,
               dossier_index=None):
    """
    tessdata : dossier des langues Tesseract (voir trouver_tessdata()) ; None = pas d'OCR.
    Renvoie (resultats, non_analyses, durees_ocr, statistiques de l'index ou None).
    """
    motif = motif_reference(numero)   # numéro invalide : erreur avant toute lecture
    documents, durees_ocr, statistiques = charger_documents(dossier, tessdata, avec_index,
                                                            reconstruire, dossier_index)
    resultats, non_analyses = chercher(motif, documents)
    return resultats, non_analyses, durees_ocr, statistiques


# ----------------------------------------------------------------------------
# Export Excel
# ----------------------------------------------------------------------------
def exporter_excel(numero, resultats, non_analyses, dossier=PDF_DIR, sortie=EXCEL_FILE):
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Résultats"

    feuille.append(["Numéro recherché", numero])
    feuille.append([])
    # « Lecture » en dernière colonne : la macro Excel lit les colonnes A à D
    feuille.append(["Fichier", "Page(s)", "Référence trouvée", "Nombres extraits", "Lecture"])
    for cellule in feuille[3]:
        cellule.font = Font(bold=True)
    feuille["A1"].font = Font(bold=True)

    for r in resultats:
        texte = " | ".join(r["nombres"])
        if len(texte) > LIMITE_CELLULE:
            texte = texte[:LIMITE_CELLULE] + " …(tronqué)"

        # Lien cliquable par formule LIEN_HYPERTEXTE (écrite en anglais dans le fichier) :
        # un lien hypertexte classique vers un chemin Windows est refusé par Excel 2007.
        # Au-delà de 255 caractères (limite d'un texte dans une formule), nom seul sans lien.
        nom = str(r["fichier"].relative_to(dossier))
        chemin = str(r["fichier"])
        if len(chemin) <= 255:
            fichier = '=HYPERLINK("{}","{}")'.format(chemin.replace('"', '""'), nom.replace('"', '""'))
        else:
            fichier = nom

        feuille.append([
            fichier,
            ", ".join(map(str, r["pages"])),
            " | ".join(r["references"]),
            texte,
            r["lecture"],
        ])
        ligne = feuille.max_row
        feuille.cell(ligne, 1).style = "Hyperlink"
        for colonne in (3, 4):
            feuille.cell(ligne, colonne).number_format = "@"
            feuille.cell(ligne, colonne).alignment = Alignment(wrap_text=True, vertical="top")

    if not resultats:
        feuille.append(["Aucun PDF ne contient ce numéro."])

    feuille.column_dimensions["A"].width = 35
    feuille.column_dimensions["B"].width = 10
    feuille.column_dimensions["C"].width = 25
    feuille.column_dimensions["D"].width = 80
    feuille.column_dimensions["E"].width = 14

    # Feuille de contrôle : ce qui n'a pas pu être lu
    controle = classeur.create_sheet("Non analysés")
    controle.append(["Fichier", "Raison"])
    for cellule in controle[1]:
        cellule.font = Font(bold=True)
    for f, raison in non_analyses:
        controle.append([str(f.relative_to(dossier)), raison])
    controle.column_dimensions["A"].width = 35
    controle.column_dimensions["B"].width = 60

    try:
        classeur.save(sortie)
    except PermissionError:
        raise PermissionError(
            f"Impossible d'écrire {sortie.name} : ferme-le dans Excel puis relance."
        )


# ----------------------------------------------------------------------------
# Ligne de commande
# ----------------------------------------------------------------------------
class AnalyseurArguments(argparse.ArgumentParser):
    """argparse sort en code 2 sur une erreur d'usage : ici, toute erreur = code 1."""

    def error(self, message):
        self.print_usage(sys.stderr)
        self.exit(1, f"Erreur : {message}\n")


def lire_arguments():
    analyseur = AnalyseurArguments(
        description="Recherche une référence dans un dossier de PDF et écrit les nombres trouvés dans Excel."
    )
    # nargs="?" : sans numéro, il est demandé au clavier (lancement par double-clic)
    analyseur.add_argument("numero", nargs="?", help="référence à rechercher, ex. COLIS-2026-00481")
    analyseur.add_argument("--dossier", type=Path, default=PDF_DIR,
                           help=f"dossier des PDF (défaut : {PDF_DIR.name} à côté de main.py)")
    analyseur.add_argument("--sortie", type=Path, default=EXCEL_FILE,
                           help=f"fichier Excel produit (défaut : {EXCEL_FILE.name} à côté de main.py)")
    analyseur.add_argument("--reindexer", action="store_true",
                           help="reconstruit l'index entièrement avant la recherche")
    analyseur.add_argument("--indexer-seulement", action="store_true",
                           help="met l'index à jour sans rechercher (numéro inutile)")
    analyseur.add_argument("--sans-index", action="store_true",
                           help="lit directement tous les PDF, sans utiliser ni modifier l'index")
    arguments = analyseur.parse_args()
    if arguments.sans_index and (arguments.reindexer or arguments.indexer_seulement):
        analyseur.error("--sans-index ne peut pas être combiné avec --reindexer ou --indexer-seulement")
    return arguments


def afficher_lecture(durees_ocr, statistiques):
    if statistiques is not None:
        print(f"Index : {statistiques['lus']} PDF lus, {statistiques['cache']} depuis le cache, "
              f"{statistiques['retires']} retirés ({statistiques['duree']:.1f} s).")
    if durees_ocr:
        print(f"OCR : {len(durees_ocr)} page(s) en {sum(durees_ocr):.1f} s.")


def dossier_documents():
    """Dossier Documents réel de l'utilisateur (il peut être redirigé, par exemple vers OneDrive)."""
    try:
        import winreg   # Windows uniquement
        cle = r"Software\Microsoft\Windows\CurrentVersion\Explorer\User Shell Folders"
        with winreg.OpenKey(winreg.HKEY_CURRENT_USER, cle) as registre:
            chemin = Path(os.path.expandvars(winreg.QueryValueEx(registre, "Personal")[0]))
        if chemin.is_dir():
            return chemin
    except (ImportError, OSError):
        pass
    return Path.home() / "Documents"


def mode_double_clic():
    """
    Exe lancé par double-clic, sans argument (plan B si la macro est bloquée) : questions dans
    la fenêtre, résultat dans Documents puis ouvert. En cas d'erreur, la fenêtre attend Entrée
    pour que le message reste lisible.
    """
    interactif = sys.stdin is not None and sys.stdin.isatty()   # faux si les saisies sont simulées
    try:
        print("Recherche dans les PDF - Diagne Digital & Tech\n")
        numero = input("Numéro à rechercher : ").lstrip("\ufeff").strip()
        reponse = input(f"Dossier des PDF (Entrée = {PDF_DIR}) : ").strip().strip('"')
        print()
        if not numero:
            print("Tu dois saisir un numéro.")
            code = 1
        else:
            dossier = Path(reponse) if reponse else PDF_DIR
            nom = re.sub(r"[^A-Za-z0-9_-]+", "_", numero)
            sortie = dossier_documents() / f"Recherche_PDF_{nom}_{time.strftime('%Y%m%d_%H%M%S')}.xlsx"
            code = executer(numero, dossier, sortie)
            if code == 0 and interactif:
                os.startfile(sortie)   # ouvre le résultat dans Excel
    except Exception as erreur:   # aucune erreur ne doit fermer la fenêtre sans message
        print(f"Erreur : {erreur}")
        code = 1
    if code != 0 and interactif:
        input("\nAppuie sur Entrée pour fermer.")
    return code


def executer(numero, dossier, sortie, indexer_seulement=False, sans_index=False, reindexer=False):
    """Recherche (ou indexation seule) avec affichage ; renvoie le code de sortie (0 ou 1)."""
    debut = time.perf_counter()

    # Chemins absolus : nécessaires pour relative_to() et pour les liens hypertextes
    dossier = Path(dossier).resolve()
    sortie = Path(sortie).resolve()

    tessdata = trouver_tessdata()
    if tessdata is None:
        print(f"Attention : Tesseract introuvable (langues {OCR_LANGUES}). "
              "Les pages scannées ne seront pas lues — voir la feuille « Non analysés ».")

    try:
        if indexer_seulement:
            _, durees_ocr, statistiques = charger_documents(dossier, tessdata, reconstruire=reindexer)
        else:
            resultats, non_analyses, durees_ocr, statistiques = rechercher(
                numero, dossier, tessdata, avec_index=not sans_index, reconstruire=reindexer)
            exporter_excel(numero, resultats, non_analyses, dossier, sortie)
    except (FileNotFoundError, PermissionError, ValueError) as e:
        print(f"Erreur : {e}")
        return 1

    afficher_lecture(durees_ocr, statistiques)
    if indexer_seulement:
        print(f"Temps total : {time.perf_counter() - debut:.1f} s.")
        return 0 if statistiques is not None else 1   # index inutilisable : il n'a pas été préparé

    print(f"{len(resultats)} PDF trouvé(s) pour « {numero} ».")
    if non_analyses:
        print(f"{len(non_analyses)} élément(s) non analysé(s) — voir la feuille « Non analysés ».")
    print(f"Résultat : {sortie}")
    print(f"Temps total : {time.perf_counter() - debut:.1f} s.")
    return 0


def main():
    if getattr(sys, "frozen", False) and len(sys.argv) == 1:
        return mode_double_clic()

    arguments = lire_arguments()
    numero = None
    if not arguments.indexer_seulement:
        numero = arguments.numero if arguments.numero is not None else input("Numéro à rechercher : ")
        numero = numero.lstrip("\ufeff").strip()   # BOM ajouté par certains terminaux (tube PowerShell)
        if not numero:
            print("Tu dois saisir un numéro.")
            return 1

    return executer(numero, arguments.dossier, arguments.sortie, arguments.indexer_seulement,
                    arguments.sans_index, arguments.reindexer)


if __name__ == "__main__":
    sys.exit(main())
